#!/usr/bin/env python3
"""Aggregate AC-balanced gate-level resource records into full Shor-map estimates.

"""

from __future__ import annotations

import argparse
import csv
import os
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Mapping

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


MODEL_VERSION = "ac-balanced-clean-v1"
FIELD_SIZES = (163, 233, 283, 571)
MAX_W_BY_N = {163: 128, 233: 128, 283: 128, 571: 256}
PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT_DIR = PROJECT_ROOT / "results"
DEFAULT_OUTPUT = DEFAULT_INPUT_DIR / "ac_balanced_full_shor.xlsx"

OUTPUT_HEADERS = [
    "n",
    "w",
    "cnot",
    "toffoli",
    "depth",
    "width",
    "toffolidepth",
    "DW",
    "TDW",
]
REQUIRED_LAYER_COLUMNS = {
    "model_version",
    "n",
    "mode",
    "lanes",
    "qubits",
    "toffoli",
    "cnot",
    "full_depth",
    "toffoli_depth",
}


@dataclass(frozen=True)
class LayerStats:
    qubits: int
    toffoli: int
    cnot: int
    full_depth: int
    toffoli_depth: int


@dataclass(frozen=True)
class FullShorEstimate:
    n: int
    w: int
    cnot: int
    toffoli: int
    depth: int
    width: int
    toffoli_depth: int

    @property
    def dw(self) -> int:
        return self.depth * self.width

    @property
    def tdw(self) -> int:
        return self.toffoli_depth * self.width

    def as_row(self) -> list[int]:
        return [
            self.n,
            self.w,
            self.cnot,
            self.toffoli,
            self.depth,
            self.width,
            self.toffoli_depth,
            self.dw,
            self.tdw,
        ]


LayerKey = tuple[str, int]


def ceil_log2(value: int) -> int:
    if value < 1:
        raise ValueError("ceil_log2 requires a positive integer")
    return (value - 1).bit_length()


def build_balanced_schedule(n: int, w: int) -> list[LayerKey]:
    """Return the exact Phase-I/Phase-II layer sequence from main.pdf Fig. 2."""
    total_inputs = 2 * n + 2
    try:
        maximum_w = MAX_W_BY_N[n]
    except KeyError as error:
        raise ValueError(f"unsupported n={n}; choose one of {FIELD_SIZES}") from error
    if not 1 <= w <= maximum_w or 2 * w > total_inputs:
        raise ValueError(
            f"Balanced schedule requires 1 <= w <= {maximum_w} for n={n} "
            "with 2w <= 2n+2"
        )

    # B0 is loaded with CNOTs.  The first cached accumulation layer adds B1.
    schedule: list[LayerKey] = [("accumulation", w)]
    remaining = total_inputs - 2 * w
    while remaining > 0:
        lanes = min(w, remaining)
        schedule.append(("accumulation", lanes))
        remaining -= lanes

    # Pair adjacent partial sums; an odd unpaired item passes to the next level.
    items = w
    while items > 1:
        lanes = items // 2
        schedule.append(("reduction", lanes))
        items = lanes + items % 2
    return schedule


def estimate_full_shor(
    n: int,
    w: int,
    layers: Mapping[LayerKey, LayerStats],
) -> FullShorEstimate:
    """
    The returned circuit is the complete controlled-point-addition map under the
    scope of main.pdf Table 6, including the outer copy/uncompute.  ``depth`` is
    the NCT full depth and ``width`` is the peak live-qubit count.
    """
    schedule = build_balanced_schedule(n, w)

    total_toffoli = 0
    total_cnot = 0
    total_full_depth = 0
    total_toffoli_depth = 0
    peak_qubits = 0
    clear_qubits = 0

    for layer_index, key in enumerate(schedule):
        try:
            stats = layers[key]
        except KeyError as error:
            raise ValueError(f"missing layer n={n}, mode={key[0]}, lanes={key[1]}") from error

        mode, lanes = key
        length = stats.qubits
        total_toffoli += stats.toffoli
        total_toffoli_depth += stats.toffoli_depth

        if layer_index == 0:
            if key != ("accumulation", w):
                raise AssertionError("invalid first Balanced layer")
            fanout_depth = 2 * ceil_log2(2 * n)
            initialization_cnot = 2 * lanes * (2 * n - 1) + 2 * n * lanes
            total_full_depth += fanout_depth + 1 + stats.full_depth
            total_cnot += initialization_cnot + stats.cnot
            peak_qubits = length - lanes
            clear_qubits = length - 6 * n * lanes - lanes
            continue

        total_full_depth += stats.full_depth
        total_cnot += stats.cnot

        if mode == "accumulation":
            demand = length - 2 * n * lanes - lanes
            if demand > clear_qubits:
                peak_qubits += demand - clear_qubits
                clear_qubits = length - 6 * n * lanes - lanes
            else:
                over = clear_qubits - demand
                clear_qubits = length - 6 * n * lanes - lanes + over
        elif mode == "reduction":
            demand = length - 4 * n * lanes
            if demand > clear_qubits:
                peak_qubits += demand - clear_qubits
                clear_qubits = length - 8 * n * lanes
            else:
                over = clear_qubits - demand
                clear_qubits = length - 8 * n * lanes + over
        else:  # pragma: no cover - schedules are constructed from two fixed modes.
            raise AssertionError(f"unknown layer mode: {mode}")

    if peak_qubits < 0:
        raise ValueError(f"negative peak width for n={n}, w={w}")

    return FullShorEstimate(
        n=n,
        w=w,
        cnot=2 * total_cnot + 2 * n,
        toffoli=2 * total_toffoli,
        depth=2 * total_full_depth + 2,
        width=peak_qubits + 4 * n + 2,
        toffoli_depth=2 * total_toffoli_depth,
    )


def _parse_nonnegative_int(raw: Mapping[str, str], column: str, path: Path, row: int) -> int:
    try:
        value = int(raw[column])
    except (KeyError, TypeError, ValueError) as error:
        raise ValueError(f"{path}: row {row}: invalid integer in {column!r}") from error
    if value < 0:
        raise ValueError(f"{path}: row {row}: {column!r} must be nonnegative")
    return value


def load_layer_cache(path: Path, n: int, maximum_w: int | None = None) -> dict[LayerKey, LayerStats]:
    """Load and validate all layers needed through ``maximum_w``."""
    configured_maximum = MAX_W_BY_N[n]
    if maximum_w is None:
        maximum_w = configured_maximum
    if not 1 <= maximum_w <= configured_maximum:
        raise ValueError(
            f"n={n}: maximum_w must satisfy 1 <= maximum_w <= {configured_maximum}"
        )
    layers: dict[LayerKey, LayerStats] = {}
    versions = set()
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        columns = set(reader.fieldnames or [])
        missing_columns = sorted(REQUIRED_LAYER_COLUMNS - columns)
        if missing_columns:
            raise ValueError(f"{path}: missing columns: {', '.join(missing_columns)}")

        for row_number, raw in enumerate(reader, start=2):
            if raw["model_version"] not in (MODEL_VERSION, "ac-balanced-clean-nct-v2"):
                raise ValueError(
                    f"{path}: row {row_number}: expected model_version={MODEL_VERSION!r}"
                )
            versions.add(raw["model_version"])
            if len(versions) > 1:
                raise ValueError(f"{path}: mixed backend versions in one cache")
            row_n = _parse_nonnegative_int(raw, "n", path, row_number)
            if row_n != n:
                raise ValueError(f"{path}: row {row_number}: expected n={n}, found n={row_n}")

            mode = raw["mode"]
            if mode not in {"accumulation", "reduction"}:
                raise ValueError(f"{path}: row {row_number}: invalid mode={mode!r}")
            lanes = _parse_nonnegative_int(raw, "lanes", path, row_number)
            if lanes < 1:
                raise ValueError(f"{path}: row {row_number}: lanes must be positive")
            key = (mode, lanes)
            if key in layers:
                raise ValueError(f"{path}: duplicate layer mode={mode}, lanes={lanes}")

            stats = LayerStats(
                qubits=_parse_nonnegative_int(raw, "qubits", path, row_number),
                toffoli=_parse_nonnegative_int(raw, "toffoli", path, row_number),
                cnot=_parse_nonnegative_int(raw, "cnot", path, row_number),
                full_depth=_parse_nonnegative_int(raw, "full_depth", path, row_number),
                toffoli_depth=_parse_nonnegative_int(
                    raw, "toffoli_depth", path, row_number
                ),
            )
            if stats.qubits == 0:
                raise ValueError(f"{path}: row {row_number}: qubits must be positive")
            layers[key] = stats

    expected = {
        *(('accumulation', lanes) for lanes in range(1, maximum_w + 1)),
        *(('reduction', lanes) for lanes in range(1, maximum_w // 2 + 1)),
    }
    missing = sorted(expected - layers.keys())
    if missing:
        preview = ", ".join(f"{mode}:{lanes}" for mode, lanes in missing[:8])
        suffix = " ..." if len(missing) > 8 else ""
        raise ValueError(f"{path}: missing {len(missing)} required layers: {preview}{suffix}")
    return layers


def resolve_layer_file(input_dir: Path, n: int) -> Path:
    """Resolve the canonical name, with a unique download-suffix fallback."""
    exact = input_dir / f"layers_n{n}.csv"
    if exact.is_file():
        return exact
    candidates = sorted(path for path in input_dir.glob(f"layers_n{n}*.csv") if path.is_file())
    if len(candidates) == 1:
        return candidates[0]
    if not candidates:
        raise FileNotFoundError(f"no layer cache found for n={n} in {input_dir}")
    names = ", ".join(path.name for path in candidates)
    raise ValueError(f"ambiguous layer caches for n={n}: {names}")


def calculate_all(
    input_dir: Path,
    maximum_w_by_n: Mapping[int, int] | None = None,
) -> tuple[dict[int, list[FullShorEstimate]], dict[int, Path]]:
    requested_limits = dict(MAX_W_BY_N if maximum_w_by_n is None else maximum_w_by_n)
    if set(requested_limits) != set(FIELD_SIZES):
        raise ValueError(f"maximum_w_by_n must define exactly {FIELD_SIZES}")
    estimates: dict[int, list[FullShorEstimate]] = {}
    sources: dict[int, Path] = {}
    for n in FIELD_SIZES:
        maximum_w = requested_limits[n]
        source = resolve_layer_file(input_dir, n)
        layers = load_layer_cache(source, n, maximum_w)
        sources[n] = source
        estimates[n] = [estimate_full_shor(n, w, layers) for w in range(1, maximum_w + 1)]
    return estimates, sources


def _style_worksheet(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{sheet.max_row}"
    widths = (8, 8, 18, 16, 16, 14, 18, 20, 20)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.number_format = "#,##0"


def export_workbook(estimates: Mapping[int, list[FullShorEstimate]], output: Path) -> None:
    workbook = Workbook()
    workbook.properties.title = "AC-based balanced Shor controlled-point-addition resources"
    workbook.properties.subject = (
        "main.pdf Table 6 scope: complete controlled point-addition map with uncompute"
    )

    for index, n in enumerate(FIELD_SIZES):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = str(n)
        sheet.append(OUTPUT_HEADERS)
        rows = estimates.get(n)
        if rows is None:
            raise ValueError(f"missing estimates for n={n}")
        widths = [row.w for row in rows]
        if not widths or widths != list(range(1, widths[-1] + 1)):
            raise ValueError(f"n={n}: w rows must be contiguous and start at 1")
        if widths[-1] > MAX_W_BY_N[n]:
            raise ValueError(f"n={n}: w exceeds configured maximum {MAX_W_BY_N[n]}")
        for estimate in rows:
            sheet.append(estimate.as_row())
        _style_worksheet(sheet)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=output.stem + ".",
            suffix=".xlsx",
            dir=output.parent,
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, output)
        output.chmod(0o644)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Aggregate complete AC-balanced layer CSV files into a four-sheet "
            "Shor controlled-point-addition resource workbook."
        )
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=DEFAULT_INPUT_DIR,
        help=f"directory containing layers_n*.csv (default: {DEFAULT_INPUT_DIR})",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"output workbook (default: {DEFAULT_OUTPUT})",
    )
    parser.add_argument(
        "--max-w-571",
        type=int,
        default=MAX_W_BY_N[571],
        metavar="W",
        help=(
            "largest n=571 width to export (default: 256); use 128 only to "
            "rebuild the currently available partial cache"
        ),
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    requested_limits = dict(MAX_W_BY_N)
    requested_limits[571] = args.max_w_571
    try:
        estimates, sources = calculate_all(args.input_dir.resolve(), requested_limits)
    except (FileNotFoundError, ValueError) as error:
        print(f"error: {error}", file=sys.stderr)
        return 2
    output = args.output.resolve()
    export_workbook(estimates, output)

    for n in FIELD_SIZES:
        minimum_dw = min(estimates[n], key=lambda row: row.dw)
        minimum_tdw = min(estimates[n], key=lambda row: row.tdw)
        print(f"n={n}: source={sources[n]}")
        print(
            f"  min DW: w={minimum_dw.w}, DW={minimum_dw.dw}; "
            f"min TDW: w={minimum_tdw.w}, TDW={minimum_tdw.tdw}"
        )
    print(f"Excel: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
