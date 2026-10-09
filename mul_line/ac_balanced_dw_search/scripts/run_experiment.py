#!/usr/bin/env python3
"""Run segmented AC-arithmetic Balanced scans and create one workbook per n."""

from __future__ import annotations

import argparse
import csv
import os
import shutil
import subprocess
import sys
import tempfile
import threading
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MODEL_VERSION = "ac-balanced-clean-nct-v2"
SUPPORTED = {163: 128, 233: 128, 283: 128, 571: 256}
PROJECT_ROOT = Path(__file__).resolve().parents[1]


def discover_repository_root(project_root: Path) -> Path:
    """Locate the submission root independently of the rebuttal layout depth."""
    override = os.environ.get("ECC_SHOR_REPOSITORY_ROOT")
    if override:
        repository_root = Path(override).expanduser().resolve()
        if (repository_root / "mul_line" / "C++" / "data").is_dir():
            return repository_root
        raise RuntimeError(
            "ECC_SHOR_REPOSITORY_ROOT does not contain mul_line/C++/data"
        )
    for candidate in (project_root, *project_root.parents):
        if (candidate / "mul_line" / "C++" / "data").is_dir():
            return candidate
    raise RuntimeError(
        "could not locate the repository root; set ECC_SHOR_REPOSITORY_ROOT"
    )


REPOSITORY_ROOT = discover_repository_root(PROJECT_ROOT)
DEFAULT_DATA_ROOT = REPOSITORY_ROOT / "mul_line" / "C++" / "data"

RESULT_HEADERS = [
    "model_version", "n", "total_inputs", "w", "accumulation_layers",
    "reduction_layers", "toffoli", "cnot", "full_depth", "current_depth",
    "toffoli_depth", "qubits", "dw_full", "dw_current", "tdw",
    "log2_dw_full", "log2_dw_current", "log2_tdw", "seconds",
]
LAYER_HEADERS = [
    "model_version", "n", "mode", "lanes", "qubits", "toffoli", "cnot",
    "full_depth", "current_depth", "toffoli_depth", "seconds", "data_root",
]

INTEGER_COLUMNS = {
    "n", "total_inputs", "w", "accumulation_layers", "reduction_layers",
    "toffoli", "cnot", "full_depth", "current_depth", "toffoli_depth", "qubits",
}
EXACT_COST_COLUMNS = {"dw_full", "dw_current", "tdw"}
FLOAT_COLUMNS = {"log2_dw_full", "log2_dw_current", "log2_tdw", "seconds"}


@dataclass(frozen=True)
class WorkerScan:
    """One independent scan worker and its private output files."""

    worker_id: int
    w_start: int
    w_end: int
    result_path: Path
    cache_path: Path


def read_latest_rows(csv_path: Path, n: int) -> list[dict[str, object]]:
    """Read the main CSV, keeping the last row when --force duplicates a w."""
    if not csv_path.exists():
        return []
    latest: dict[int, dict[str, object]] = {}
    with csv_path.open(newline="", encoding="utf-8") as handle:
        for raw in csv.DictReader(handle):
            if raw.get("model_version") != MODEL_VERSION or int(raw["n"]) != n:
                continue
            row: dict[str, object] = dict(raw)
            for key in INTEGER_COLUMNS:
                row[key] = int(raw[key])
            for key in EXACT_COST_COLUMNS:
                row[key] = raw[key]
            for key in FLOAT_COLUMNS:
                row[key] = float(raw[key])
            latest[int(raw["w"])] = row
    return [latest[w] for w in sorted(latest)]


def contiguous_ranges(values: Iterable[int]) -> str:
    numbers = sorted(set(values))
    if not numbers:
        return "(none)"
    ranges: list[str] = []
    start = previous = numbers[0]
    for value in numbers[1:]:
        if value == previous + 1:
            previous = value
            continue
        ranges.append(str(start) if start == previous else f"{start}-{previous}")
        start = previous = value
    ranges.append(str(start) if start == previous else f"{start}-{previous}")
    return ", ".join(ranges)


def partition_pending_widths(widths: Iterable[int], jobs: int) -> list[tuple[int, int]]:
    if jobs < 1:
        raise ValueError("jobs must be at least 1")
    pending = sorted(set(widths))
    if not pending:
        return []

    worker_count = min(jobs, len(pending))
    quotient, remainder = divmod(len(pending), worker_count)
    chunks: list[tuple[int, int]] = []
    offset = 0
    for worker_index in range(worker_count):
        size = quotient + (1 if worker_index < remainder else 0)
        group = pending[offset:offset + size]
        chunks.append((group[0], group[-1]))
        offset += size
    return chunks


def scan_widths(w_start: int, w_end: int, gap: int) -> range:
    """Return the w sequence starting at start, stepping by gap, through end."""
    if gap < 1:
        raise ValueError("g must be at least 1")
    return range(w_start, w_end + 1, gap)


def read_raw_csv(path: Path) -> list[dict[str, str]]:
    """Read raw CSV rows, ignoring blanks and truncated rows with missing columns."""
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as handle:
        rows: list[dict[str, str]] = []
        for raw in csv.DictReader(handle):
            if not raw or None in raw or any(value is None for value in raw.values()):
                continue
            rows.append(dict(raw))
        return rows


def write_csv_atomic(
    path: Path, headers: list[str], rows: Iterable[dict[str, object]],
) -> None:
    """Write a sibling temporary file and atomically replace the destination CSV."""
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w", newline="", encoding="utf-8", prefix=path.stem + ".",
            suffix=".csv", dir=path.parent, delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
            writer = csv.DictWriter(
                temporary, fieldnames=headers, extrasaction="ignore", lineterminator="\n",
            )
            writer.writeheader()
            for row in rows:
                writer.writerow({header: row.get(header, "") for header in headers})
        os.replace(temporary_path, path)
        path.chmod(0o644)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def merge_parallel_results(n: int, csv_path: Path, workers: Iterable[WorkerScan]) -> None:
    """Merge completed worker results, keeping a worker's last row for the same w."""
    passthrough: list[dict[str, str]] = []
    latest: dict[int, dict[str, str]] = {}

    for raw in read_raw_csv(csv_path):
        try:
            is_current = raw.get("model_version") == MODEL_VERSION and int(raw["n"]) == n
            width = int(raw["w"])
        except (KeyError, TypeError, ValueError):
            is_current = False
            width = -1
        if is_current:
            latest[width] = raw
        else:
            passthrough.append(raw)

    for worker in workers:
        for raw in read_raw_csv(worker.result_path):
            try:
                if raw.get("model_version") != MODEL_VERSION or int(raw["n"]) != n:
                    continue
                width = int(raw["w"])
            except (KeyError, TypeError, ValueError):
                continue
            if worker.w_start <= width <= worker.w_end:
                latest[width] = raw

    merged = passthrough + [latest[width] for width in sorted(latest)]
    write_csv_atomic(csv_path, RESULT_HEADERS, merged)


def merge_parallel_caches(n: int, cache_path: Path, workers: Iterable[WorkerScan]) -> None:
    """Deduplicate by (mode, lanes) and atomically merge relocatable caches."""
    passthrough: list[dict[str, str]] = []
    latest: dict[tuple[str, int], dict[str, str]] = {}

    def ingest(raw: dict[str, str], preserve_unrelated: bool) -> None:
        try:
            is_current = raw.get("model_version") == MODEL_VERSION and int(raw["n"]) == n
            key = (raw["mode"], int(raw["lanes"]))
        except (KeyError, TypeError, ValueError):
            is_current = False
            key = ("", -1)
        if is_current:
            latest[key] = raw
        elif preserve_unrelated:
            passthrough.append(raw)

    for raw in read_raw_csv(cache_path):
        ingest(raw, preserve_unrelated=True)
    for worker in workers:
        for raw in read_raw_csv(worker.cache_path):
            ingest(raw, preserve_unrelated=False)

    ordered = [latest[key] for key in sorted(latest, key=lambda item: (item[0], item[1]))]
    write_csv_atomic(cache_path, LAYER_HEADERS, passthrough + ordered)


def read_layer_rows(cache_path: Path, n: int) -> list[dict[str, object]]:
    if not cache_path.exists():
        return []
    latest: dict[tuple[str, int], dict[str, object]] = {}
    with cache_path.open(newline="", encoding="utf-8") as handle:
        for raw in csv.DictReader(handle):
            if raw.get("model_version") != MODEL_VERSION or int(raw["n"]) != n:
                continue
            key = (raw["mode"], int(raw["lanes"]))
            latest[key] = {
                "mode": raw["mode"],
                "lanes": int(raw["lanes"]),
                "qubits": int(raw["qubits"]),
                "toffoli": int(raw["toffoli"]),
                "cnot": int(raw["cnot"]),
                "full_depth": int(raw["full_depth"]),
                "current_depth": int(raw["current_depth"]),
                "toffoli_depth": int(raw["toffoli_depth"]),
                "seconds": float(raw["seconds"]),
                "data_root": raw["data_root"],
            }
    return [latest[key] for key in sorted(latest, key=lambda item: (item[0], item[1]))]


def style_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(width + 2, 28)


def export_workbook(n: int, csv_path: Path, cache_path: Path, workbook_path: Path) -> None:
    rows = read_latest_rows(csv_path, n)
    layers = read_layer_rows(cache_path, n)

    workbook = Workbook()
    results = workbook.active
    results.title = "Results"
    results.append(RESULT_HEADERS)
    for row in rows:
        results.append([row[column] for column in RESULT_HEADERS])
    for column in (13, 14, 15):
        for row_number in range(2, results.max_row + 1):
            results.cell(row=row_number, column=column).number_format = "@"
    style_sheet(results)

    summary = workbook.create_sheet("Summary")
    summary.append(["Item", "Value"])
    summary.append(["model_version", MODEL_VERSION])
    summary.append(["n", n])
    summary.append(["Shor controlled inputs N", 2 * n + 2])
    summary.append(["Allowed w", f"1-{SUPPORTED[n]}"])
    summary.append(["Currently scanned ranges", contiguous_ranges(int(row["w"]) for row in rows)])
    summary.append(["Completed w count", len(rows)])
    if rows:
        min_dw = min(rows, key=lambda row: int(str(row["dw_full"])))
        min_tdw = min(rows, key=lambda row: int(str(row["tdw"])))
        summary.append(["arg min DW(w)", int(min_dw["w"])])
        summary.append(["min DW(w)=full_depth×qubits", str(min_dw["dw_full"])])
        summary.append(["At this point (full_depth, qubits)", f"({min_dw['full_depth']}, {min_dw['qubits']})"])
        summary.append(["arg min TDW(w)", int(min_tdw["w"])])
        summary.append(["min TDW(w)=toffoli_depth×qubits", str(min_tdw["tdw"])])
        summary.append(["At this point (toffoli_depth, qubits)", f"({min_tdw['toffoli_depth']}, {min_tdw['qubits']})"])
    summary.append(["DW definition", "NCT full depth x width, as in paper Table 6"])
    summary.append(["Point-addition structure", "Balanced clean-copy/uncompute, not T-optimal retained output"])
    summary.append(["Arithmetic", "AC multiplication matrices + AC Itoh–Tsujii inversion chain"])
    style_sheet(summary)
    summary.auto_filter.ref = "A1:B1"

    cache_sheet = workbook.create_sheet("LayerCache")
    layer_export_headers = [
        "mode", "lanes", "qubits", "toffoli", "cnot", "full_depth",
        "current_depth", "toffoli_depth", "seconds", "data_root",
    ]
    cache_sheet.append(layer_export_headers)
    for row in layers:
        cache_sheet.append([row[column] for column in layer_export_headers])
    style_sheet(cache_sheet)

    if n == 163 and rows:
        optimum = min(rows, key=lambda row: int(str(row["dw_full"])))
        comparison = workbook.create_sheet("Comparison")
        comparison_headers = [
            "design", "point_add_structure", "w", "toffoli", "cnot", "full_depth",
            "toffoli_depth", "qubits", "dw_full", "tdw",
        ]
        comparison.append(comparison_headers)
        comparison.append([
            "Table 6 Karatsuba Balanced", "Balanced clean-copy/uncompute", 30,
            30464596, 265802254, 73788, 1104, 873913,
            "64484292444", "964799952",
        ])
        comparison.append([
            "Table 6 AC T-optimal", "T-optimal retained-output", 163,
            3167396, 571478060, 233752, 322, 1643885,
            "384261406520", "529330970",
        ])
        comparison.append([
            "This project AC Balanced", "Balanced clean-copy/uncompute", int(optimum["w"]),
            int(optimum["toffoli"]), int(optimum["cnot"]), int(optimum["full_depth"]),
            int(optimum["toffoli_depth"]), int(optimum["qubits"]),
            str(optimum["dw_full"]), str(optimum["tdw"]),
        ])
        for column in (9, 10):
            for row_number in range(2, comparison.max_row + 1):
                comparison.cell(row=row_number, column=column).number_format = "@"
        style_sheet(comparison)

    assumptions = workbook.create_sheet("Assumptions")
    assumptions.append(["No.", "Description"])
    assumptions.append([1, "N=2n+2; a controlled write initializes B0, followed by ceil(N/w)-1 controlled point-addition layers."])
    assumptions.append([2, "Phase II performs a binary-tree reduction of the w partial sums; each layer uses the Montgomery trick for parallel inversion."])
    assumptions.append([3, "Inverse and multiplication results are copied to dedicated n-bit outputs before reversing the stream, allowing reuse of AC multiplication ancillas."])
    assumptions.append([4, "As in the original C++ estimator, the complete Shor map includes compute/uncompute, so gate counts and depths are doubled while width is not."])
    assumptions.append([5, "Exact DW/TDW values are stored as text in Excel; log2 columns support plots and numerical comparisons."])
    style_sheet(assumptions)

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        prefix=workbook_path.stem + ".", suffix=".xlsx", dir=workbook_path.parent,
        delete=False,
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
        workbook_path.chmod(0o644)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Scan w for AC-based arithmetic with the Balanced strategy, in segments or parallel workers, and update Excel."
    )
    parser.add_argument("--n", type=int, required=True, choices=sorted(SUPPORTED))
    parser.add_argument("--w-start", type=int, help="first w in this scan, inclusive")
    parser.add_argument("--w-end", type=int, help="last w in this scan, inclusive")
    parser.add_argument("--data-root", type=Path, default=DEFAULT_DATA_ROOT)
    parser.add_argument("--results-dir", type=Path, default=PROJECT_ROOT / "results")
    parser.add_argument("--cache-dir", type=Path, default=PROJECT_ROOT / "cache")
    parser.add_argument("--binary", type=Path, default=PROJECT_ROOT / "build" / "ac_balanced_estimator")
    parser.add_argument(
        "--jobs", type=int, default=1, metavar="N",
        help="independent scan workers to run concurrently (default: 1; each consumes one peak allocation)",
    )
    parser.add_argument(
        "--g", type=int, default=1, metavar="G",
        help="w stride (default: 1); for example, start=1, end=16, G=3 scans 1,4,...,16",
    )
    parser.add_argument("--no-build", action="store_true")
    parser.add_argument("--force", action="store_true", help="re-estimate w even when it already exists")
    parser.add_argument("--export-only", action="store_true", help="rebuild Excel from the existing CSV without running C++")
    return parser


def scan_command(
    args: argparse.Namespace, w_start: int, w_end: int,
    result_path: Path, cache_path: Path,
) -> list[str]:
    command = [
        str(args.binary), "scan", "--n", str(args.n),
        "--w-start", str(w_start), "--w-end", str(w_end),
        "--g", str(args.g),
        "--data-root", str(args.data_root.resolve()),
        "--output-csv", str(result_path.resolve()),
        "--layer-cache", str(cache_path.resolve()),
    ]
    if args.force:
        command.append("--force")
    return command


def run_serial_scan(args: argparse.Namespace, csv_path: Path, cache_path: Path) -> int:
    """Preserve jobs=1 behavior: one C++ process scans by g and writes immediately."""
    command = scan_command(args, args.w_start, args.w_end, csv_path, cache_path)
    try:
        subprocess.run(command, cwd=PROJECT_ROOT, check=True)
    except subprocess.CalledProcessError as error:
        return error.returncode
    return 0


def relay_worker_output(worker_id: int, process: subprocess.Popen[str]) -> None:
    """Forward one worker's output linewise to prevent interleaving between processes."""
    if process.stdout is None:
        return
    for line in process.stdout:
        print(f"[worker {worker_id}] {line}", end="", flush=True)
    process.stdout.close()


def run_parallel_scan(args: argparse.Namespace, csv_path: Path, cache_path: Path) -> int:
    """Run isolated C++ workers and safely merge every successfully flushed result."""
    completed = {int(row["w"]) for row in read_latest_rows(csv_path, args.n)}
    requested = scan_widths(args.w_start, args.w_end, args.g)
    pending = list(requested) if args.force else [width for width in requested if width not in completed]
    chunks = partition_pending_widths(pending, args.jobs)
    if not chunks:
        print("[parallel] No pending w in the selected interval; skipping the C++ scan.", flush=True)
        return 0

    print(
        f"[parallel] requested jobs={args.jobs}, active workers={len(chunks)}, "
        f"g={args.g}, pending w={len(pending)}",
        flush=True,
    )
    print(
        "[parallel] chunks=" + ", ".join(f"{start}-{end}" for start, end in chunks),
        flush=True,
    )

    with tempfile.TemporaryDirectory(
        prefix=f".parallel_n{args.n}_", dir=args.cache_dir.resolve(),
    ) as temporary_directory:
        worker_root = Path(temporary_directory)
        print(f"[parallel] private worker root={worker_root}", flush=True)
        workers: list[WorkerScan] = []
        for worker_id, (w_start, w_end) in enumerate(chunks, start=1):
            private_directory = worker_root / f"worker_{worker_id}"
            private_directory.mkdir(parents=True)
            result_path = private_directory / csv_path.name
            private_cache_path = private_directory / cache_path.name
            if csv_path.exists():
                shutil.copy2(csv_path, result_path)
            if cache_path.exists():
                shutil.copy2(cache_path, private_cache_path)
            workers.append(WorkerScan(
                worker_id=worker_id,
                w_start=w_start,
                w_end=w_end,
                result_path=result_path,
                cache_path=private_cache_path,
            ))

        processes: list[tuple[WorkerScan, subprocess.Popen[str], threading.Thread]] = []
        statuses: list[tuple[WorkerScan, int]] = []
        try:
            try:
                for worker in workers:
                    print(
                        f"[parallel] launch worker={worker.worker_id} "
                        f"w={worker.w_start}-{worker.w_end}",
                        flush=True,
                    )
                    process = subprocess.Popen(
                        scan_command(
                            args, worker.w_start, worker.w_end,
                            worker.result_path, worker.cache_path,
                        ),
                        cwd=PROJECT_ROOT,
                        stdout=subprocess.PIPE,
                        stderr=subprocess.STDOUT,
                        text=True,
                        bufsize=1,
                    )
                    relay = threading.Thread(
                        target=relay_worker_output,
                        args=(worker.worker_id, process),
                        name=f"scan-output-{worker.worker_id}",
                        daemon=True,
                    )
                    relay.start()
                    processes.append((worker, process, relay))
                for worker, process, relay in processes:
                    return_code = process.wait()
                    relay.join()
                    statuses.append((worker, return_code))
                    print(
                        f"[parallel] finish worker={worker.worker_id} "
                        f"w={worker.w_start}-{worker.w_end} returncode={return_code}",
                        flush=True,
                    )
            except BaseException:
                for _, process, _ in processes:
                    if process.poll() is None:
                        process.terminate()
                for _, process, relay in processes:
                    if process.poll() is None:
                        process.wait()
                    relay.join()
                raise
        finally:
            merge_parallel_results(args.n, csv_path, workers)
            merge_parallel_caches(args.n, cache_path, workers)

    failures = [(worker, code) for worker, code in statuses if code != 0]
    if failures:
        for worker, code in failures:
            print(
                f"[parallel] worker={worker.worker_id} w={worker.w_start}-{worker.w_end} "
                f"failed (returncode={code}); results completed before failure were merged.",
                file=sys.stderr,
            )
        return 1
    return 0


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    maximum = SUPPORTED[args.n]
    if args.jobs < 1:
        raise SystemExit("--jobs N requires N >= 1")
    if args.g < 1:
        raise SystemExit("--g G requires G >= 1")
    if not args.export_only:
        if args.w_start is None or args.w_end is None:
            raise SystemExit("--w-start and --w-end are required unless --export-only is used")
        if not (1 <= args.w_start <= args.w_end <= maximum):
            raise SystemExit(f"n={args.n} requires 1 <= w-start <= w-end <= {maximum}")

    args.results_dir.mkdir(parents=True, exist_ok=True)
    args.cache_dir.mkdir(parents=True, exist_ok=True)
    csv_path = args.results_dir / f"ac_balanced_n{args.n}.csv"
    workbook_path = args.results_dir / f"ac_balanced_n{args.n}.xlsx"
    cache_path = args.cache_dir / f"layers_n{args.n}.csv"

    return_code = 0
    if not args.export_only:
        if not args.no_build:
            subprocess.run(["make", "-C", str(PROJECT_ROOT)], check=True)
        try:
            if args.jobs == 1:
                return_code = run_serial_scan(args, csv_path, cache_path)
            else:
                return_code = run_parallel_scan(args, csv_path, cache_path)
        finally:
            export_workbook(args.n, csv_path, cache_path, workbook_path)
    else:
        export_workbook(args.n, csv_path, cache_path, workbook_path)

    rows = read_latest_rows(csv_path, args.n)
    print(f"Excel: {workbook_path}")
    print(f"Completed w: {contiguous_ranges(int(row['w']) for row in rows)}")
    if rows:
        min_dw = min(rows, key=lambda row: int(str(row["dw_full"])))
        min_tdw = min(rows, key=lambda row: int(str(row["tdw"])))
        print(f"Current min DW: w={min_dw['w']}, DW={min_dw['dw_full']}")
        print(f"Current min TDW: w={min_tdw['w']}, TDW={min_tdw['tdw']}")
    return return_code


if __name__ == "__main__":
    sys.exit(main())
