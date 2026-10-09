import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.run_experiment import (
    LAYER_HEADERS,
    MODEL_VERSION,
    RESULT_HEADERS,
    SUPPORTED,
    WorkerScan,
    build_parser,
    contiguous_ranges,
    export_workbook,
    merge_parallel_caches,
    merge_parallel_results,
    partition_pending_widths,
    read_latest_rows,
    scan_command,
    scan_widths,
)


class ExportTests(unittest.TestCase):
    def test_n571_supports_w_up_to_256(self):
        self.assertEqual(SUPPORTED[571], 256)

    def test_contiguous_ranges(self):
        self.assertEqual(contiguous_ranges([1, 2, 3, 7, 9, 10]), "1-3, 7, 9-10")

    def test_partition_pending_widths(self):
        self.assertEqual(
            partition_pending_widths(range(1, 11), 3),
            [(1, 4), (5, 7), (8, 10)],
        )
        self.assertEqual(
            partition_pending_widths([1, 2, 5, 9, 10], 3),
            [(1, 2), (5, 9), (10, 10)],
        )
        self.assertEqual(partition_pending_widths([7, 8], 20), [(7, 7), (8, 8)])
        self.assertEqual(
            partition_pending_widths([1, 4, 7, 10, 13, 16], 2),
            [(1, 7), (10, 16)],
        )
        self.assertEqual(partition_pending_widths([], 4), [])
        with self.assertRaises(ValueError):
            partition_pending_widths([1], 0)

    def test_scan_widths_uses_inclusive_end_and_gap(self):
        self.assertEqual(list(scan_widths(1, 16, 3)), [1, 4, 7, 10, 13, 16])
        self.assertEqual(list(scan_widths(2, 16, 5)), [2, 7, 12])
        with self.assertRaises(ValueError):
            scan_widths(1, 16, 0)

    def test_gap_argument_is_forwarded_to_cpp_scan(self):
        args = build_parser().parse_args([
            "--n", "233", "--w-start", "1", "--w-end", "16",
            "--g", "3", "--jobs", "2",
        ])
        command = scan_command(
            args, 1, 7, Path("/tmp/results.csv"), Path("/tmp/layers.csv"),
        )
        self.assertEqual(args.g, 3)
        self.assertEqual(command[command.index("--g") + 1], "3")

    def test_latest_row_and_workbook_minima(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            csv_path = root / "result.csv"
            cache_path = root / "cache.csv"
            workbook_path = root / "result.xlsx"
            headers = [
                *RESULT_HEADERS,
            ]
            rows = [
                [MODEL_VERSION, 163, 328, 1, 327, 0, 10, 20, 100, 110, 8, 1000,
                 "100000", "110000", "8000", 16.6, 16.7, 12.9, 1.0],
                [MODEL_VERSION, 163, 328, 2, 163, 1, 11, 21, 80, 90, 6, 1100,
                 "88000", "99000", "6600", 16.4, 16.6, 12.7, 2.0],
                # Duplicate w=2; the final row must replace the old value.
                [MODEL_VERSION, 163, 328, 2, 163, 1, 12, 22, 70, 85, 7, 1100,
                 "77000", "93500", "7700", 16.2, 16.5, 12.9, 3.0],
            ]
            with csv_path.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow(headers)
                writer.writerows(rows)

            latest = read_latest_rows(csv_path, 163)
            self.assertEqual(len(latest), 2)
            self.assertEqual(latest[1]["dw_full"], "77000")

            export_workbook(163, csv_path, cache_path, workbook_path)
            workbook = load_workbook(workbook_path, data_only=True)
            summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(summary["arg min DW(w)"], 2)
            self.assertEqual(summary["arg min TDW(w)"], 2)
            self.assertEqual(summary["Currently scanned ranges"], "1-2")
            self.assertIn("Comparison", workbook.sheetnames)

    def test_empty_parameter_workbook(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_path = root / "empty.xlsx"
            export_workbook(571, root / "missing.csv", root / "missing-cache.csv", workbook_path)
            workbook = load_workbook(workbook_path, data_only=True)
            summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(summary["Currently scanned ranges"], "(none)")
            self.assertEqual(summary["Completed w count"], 0)

    def test_parallel_result_merge_keeps_worker_owned_latest_rows(self):
        def result_row(width: int, toffoli: int) -> list[object]:
            return [
                MODEL_VERSION, 163, 328, width, 1, 1, toffoli, 20, 100, 110,
                8, 1000, "100000", "110000", "8000", 16.6, 16.7, 12.9, 1.0,
            ]

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            canonical = root / "result.csv"
            worker1_path = root / "worker1.csv"
            worker2_path = root / "worker2.csv"
            for path, rows in (
                (canonical, [result_row(1, 10), result_row(2, 20)]),
                # Both files contain the main CSV seed; the last w=1 row is the --force result.
                (worker1_path, [result_row(1, 10), result_row(2, 20), result_row(1, 101)]),
                (worker2_path, [result_row(1, 10), result_row(2, 20), result_row(3, 303)]),
            ):
                with path.open("w", newline="", encoding="utf-8") as handle:
                    writer = csv.writer(handle)
                    writer.writerow(RESULT_HEADERS)
                    writer.writerows(rows)

            workers = [
                WorkerScan(1, 1, 1, worker1_path, root / "cache1.csv"),
                WorkerScan(2, 3, 3, worker2_path, root / "cache2.csv"),
            ]
            merge_parallel_results(163, canonical, workers)
            latest = {int(row["w"]): row for row in read_latest_rows(canonical, 163)}
            self.assertEqual(sorted(latest), [1, 2, 3])
            self.assertEqual(latest[1]["toffoli"], 101)
            self.assertEqual(latest[2]["toffoli"], 20)
            self.assertEqual(latest[3]["toffoli"], 303)

    def test_parallel_cache_merge_deduplicates_layer_keys(self):
        def cache_row(mode: str, lanes: int, toffoli: int) -> list[object]:
            return [
                MODEL_VERSION, 163, mode, lanes, 1000, toffoli, 20, 30, 31, 2,
                1.0, "/data",
            ]

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            canonical = root / "cache.csv"
            worker1_path = root / "worker1-cache.csv"
            worker2_path = root / "worker2-cache.csv"
            for path, rows in (
                (canonical, [cache_row("accumulation", 1, 10)]),
                (worker1_path, [
                    cache_row("accumulation", 1, 10),
                    cache_row("accumulation", 2, 22),
                ]),
                (worker2_path, [
                    cache_row("accumulation", 1, 10),
                    cache_row("reduction", 1, 33),
                ]),
            ):
                with path.open("w", newline="", encoding="utf-8") as handle:
                    writer = csv.writer(handle)
                    writer.writerow(LAYER_HEADERS)
                    writer.writerows(rows)

            workers = [
                WorkerScan(1, 1, 1, root / "result1.csv", worker1_path),
                WorkerScan(2, 2, 2, root / "result2.csv", worker2_path),
            ]
            merge_parallel_caches(163, canonical, workers)
            with canonical.open(newline="", encoding="utf-8") as handle:
                rows = list(csv.DictReader(handle))
            keys = {(row["mode"], int(row["lanes"])) for row in rows}
            self.assertEqual(
                keys,
                {("accumulation", 1), ("accumulation", 2), ("reduction", 1)},
            )
            self.assertEqual(len(rows), 3)
            self.assertNotIn(b"\r\n", canonical.read_bytes())


if __name__ == "__main__":
    unittest.main()
