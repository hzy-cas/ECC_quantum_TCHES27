import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.export_full_shor import (
    FIELD_SIZES,
    MAX_W_BY_N,
    OUTPUT_HEADERS,
    LayerStats,
    build_balanced_schedule,
    calculate_all,
    estimate_full_shor,
    export_workbook,
)


ROOT = Path(__file__).resolve().parents[1]
AVAILABLE_MAX_W_BY_N = {**MAX_W_BY_N, 571: 128}


class FullShorExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.estimates, cls.sources = calculate_all(ROOT / "results", AVAILABLE_MAX_W_BY_N)

    def test_schedule_matches_balanced_cpp_regression(self):
        schedule = build_balanced_schedule(163, 30)
        accumulation = [lanes for mode, lanes in schedule if mode == "accumulation"]
        reduction = [lanes for mode, lanes in schedule if mode == "reduction"]
        self.assertEqual(accumulation, [30, 30, 30, 30, 30, 30, 30, 30, 30, 28])
        self.assertEqual(reduction, [15, 7, 4, 2, 1])

    def test_n571_schedule_reaches_w256(self):
        schedule = build_balanced_schedule(571, 256)
        accumulation = [lanes for mode, lanes in schedule if mode == "accumulation"]
        reduction = [lanes for mode, lanes in schedule if mode == "reduction"]
        self.assertEqual(accumulation, [256, 256, 256, 120])
        self.assertEqual(reduction, [128, 64, 32, 16, 8, 4, 2, 1])

    def test_python_aggregation_matches_cpp_fake_layer_regression(self):
        fake = LayerStats(
            qubits=10000,
            toffoli=2,
            cnot=3,
            full_depth=5,
            toffoli_depth=11,
        )
        result = estimate_full_shor(163, 1, {("accumulation", 1): fake})
        self.assertEqual(result.toffoli, 1308)
        self.assertEqual(result.cnot, 4240)
        self.assertEqual(result.depth, 3310)
        self.assertEqual(result.toffoli_depth, 7194)
        self.assertEqual(result.width, 223205)
        self.assertEqual(result.dw, 3310 * 223205)
        self.assertEqual(result.tdw, 7194 * 223205)

    def test_n163_known_complete_result(self):
        result = self.estimates[163][55]
        self.assertEqual(result.w, 56)
        self.assertEqual(result.cnot, 857859900)
        self.assertEqual(result.toffoli, 6341768)
        self.assertEqual(result.depth, 577596)
        self.assertEqual(result.width, 512450)
        self.assertEqual(result.toffoli_depth, 810)
        self.assertEqual(result.dw, 295989070200)
        self.assertEqual(result.tdw, 415084500)

    def test_all_inputs_and_results_are_complete(self):
        self.assertEqual(set(self.sources), set(FIELD_SIZES))
        for n in FIELD_SIZES:
            with self.subTest(n=n):
                self.assertTrue(self.sources[n].is_file())
                maximum_w = AVAILABLE_MAX_W_BY_N[n]
                self.assertEqual([row.w for row in self.estimates[n]], list(range(1, maximum_w + 1)))
                for row in self.estimates[n]:
                    self.assertEqual(row.dw, row.depth * row.width)
                    self.assertEqual(row.tdw, row.toffoli_depth * row.width)

    def test_workbook_has_four_exact_tables(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "full_shor.xlsx"
            export_workbook(self.estimates, output)
            workbook = load_workbook(output, data_only=True, read_only=True)
            self.assertEqual(workbook.sheetnames, [str(n) for n in FIELD_SIZES])
            for n in FIELD_SIZES:
                with self.subTest(n=n):
                    sheet = workbook[str(n)]
                    self.assertEqual(sheet.max_row, AVAILABLE_MAX_W_BY_N[n] + 1)
                    self.assertEqual(sheet.max_column, len(OUTPUT_HEADERS))
                    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
                    self.assertEqual(headers, OUTPUT_HEADERS)
                    for row in sheet.iter_rows(min_row=2, values_only=True):
                        values = dict(zip(OUTPUT_HEADERS, row))
                        self.assertEqual(values["n"], n)
                        self.assertEqual(values["DW"], values["depth"] * values["width"])
                        self.assertEqual(
                            values["TDW"], values["toffolidepth"] * values["width"]
                        )


if __name__ == "__main__":
    unittest.main()
